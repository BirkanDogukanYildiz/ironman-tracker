# -*- coding: utf-8 -*-
"""Veritabanındaki kayıtlardan, formülleri canlı 9 sayfalık .xlsx üretir.

Üretilen dosya, uygulamadan bağımsız olarak Excel'de çalışmaya devam eder:
tüm özetler, seviyeler ve hedefler gerçek Excel formülleriyle hesaplanır.
"""
from __future__ import annotations

import datetime as dt
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, PieChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.marker import Marker
from openpyxl.chart.series import DataPoint
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.comments import Comment
from openpyxl.drawing.line import LineProperties
from openpyxl.formatting.rule import CellIsRule, DataBarRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.properties import PageSetupProperties

from .config import LEVEL_DETAIL, LEVEL_NAMES, SPORT_META, SPORT_ORDER, WORKOUT_KINDS

# --------------------------------------------------------------------------
# Stil sabitleri
# --------------------------------------------------------------------------
FONT = "Arial"
NAVY, SLATE, INK, MUTED = "0B1220", "1E293B", "0F172A", "64748B"
LINE, BG, WHITE, HEADBG = "E2E8F0", "F8FAFC", "FFFFFF", "1E293B"
SWIM, SWIM_L, SWIM_XL = "0284C7", "E0F2FE", "F0F9FF"
BIKE, BIKE_L, BIKE_XL = "D97706", "FEF3C7", "FFFBEB"
RUN, RUN_L, RUN_XL = "059669", "D1FAE5", "ECFDF5"
IRON, IRON_L = "E11D48", "FFE4E6"
OK_G, OK_BG, GREY_BG = "16A34A", "DCFCE7", "F1F5F9"
INPUT_FILL, INPUT_FONT = "FFF8E1", "1D4ED8"

L_SWIM, L_BIKE, L_RUN = "🏊 Yüzme", "🚴 Bisiklet", "🏃 Koşu"
SPORT_LABEL = {"swim": L_SWIM, "bike": L_BIKE, "run": L_RUN}

FMT_DATE, FMT_KM2, FMT_KM1, FMT_INT = "dd.mm.yyyy", "#,##0.00", "#,##0.0", "#,##0"
FMT_PACE, FMT_SPD, FMT_PCT = "mm:ss", "0.0", "0.0%"
FMT_HR = '#,##0.0" saat"'
FMT_KMU2, FMT_KMU1 = '#,##0.00" km"', '#,##0.0" km"'
FMT_ADET = '#,##0" antrenman"'
FMT_DPCT = '+0.0%;-0.0%;0.0%'

SH_DASH, SH_LVL, SH_LOG = "DASHBOARD", "SEVİYELER", "ANTRENMAN KAYIT"
SH_WEEK, SH_BR, SH_BRICK = "HAFTALIK TAKİP", "BRANŞ İLERLEME", "BRICK"
SH_GOAL, SH_HELP, SH_SET = "HEDEFLER", "NASIL KULLANILIR", "AYARLAR"
BRQ = f"'{SH_BR}'"


# --------------------------------------------------------------------------
# Küçük yardımcılar
# --------------------------------------------------------------------------
def f(size=10, bold=False, color=INK, italic=False):
    return Font(name=FONT, size=size, bold=bold, color=color, italic=italic)


def fill(hexcolor):
    return PatternFill("solid", fgColor=hexcolor)


def al(h="left", v="center", wrap=False, indent=0):
    return Alignment(horizontal=h, vertical=v, wrap_text=wrap, indent=indent)


def side(color, style="thin"):
    return Side(style=style, color=color)


THIN = Side(style="thin", color=LINE)
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def apply_box(ws, r1, c1, r2, c2, color=LINE, left_accent=None):
    s = side(color)
    for r in range(r1, r2 + 1):
        for c in range(c1, c2 + 1):
            cell = ws.cell(row=r, column=c)
            b = cell.border
            left, right, top, bottom = b.left, b.right, b.top, b.bottom
            if c == c1:
                left = side(left_accent, "medium") if left_accent else s
            if c == c2:
                right = s
            if r == r1:
                top = s
            if r == r2:
                bottom = s
            cell.border = Border(left=left, right=right, top=top, bottom=bottom)


def fill_range(ws, r1, c1, r2, c2, hexcolor):
    pf = fill(hexcolor)
    for r in range(r1, r2 + 1):
        for c in range(c1, c2 + 1):
            ws.cell(row=r, column=c).fill = pf


def merge_set(ws, r1, c1, r2, c2, value=None, font=None, alignment=None,
              fillc=None, numfmt=None):
    ws.merge_cells(start_row=r1, start_column=c1, end_row=r2, end_column=c2)
    cell = ws.cell(row=r1, column=c1)
    if value is not None:
        cell.value = value
    if font is not None:
        cell.font = font
    if alignment is not None:
        cell.alignment = alignment
    if numfmt is not None:
        cell.number_format = numfmt
    if fillc is not None:
        fill_range(ws, r1, c1, r2, c2, fillc)
    return cell


def widths(ws, mapping):
    for col, w in mapping.items():
        ws.column_dimensions[col].width = w


def header_row(ws, row, headers, start_col=1, height=30, bg=HEADBG, size=9.5):
    ws.row_dimensions[row].height = height
    for i, h in enumerate(headers):
        c = ws.cell(row=row, column=start_col + i, value=h)
        c.font = f(size, True, "FFFFFF")
        c.fill = fill(bg)
        c.alignment = al("center", "center", True)
        c.border = Border(left=side("334155"), right=side("334155"),
                          top=side("334155"), bottom=side("334155"))


def datestr(ref):
    return f'DAY({ref})&"."&MONTH({ref})&"."&YEAR({ref})'


def _to_date(value):
    if isinstance(value, dt.date):
        return value
    if not value:
        return None
    try:
        return dt.datetime.strptime(str(value)[:10], "%Y-%m-%d").date()
    except ValueError:
        return None


# ==========================================================================
def build_workbook(out_path, *, workouts=None, goals=None, thresholds=None,
                   settings=None, n_week=160):
    """Çalışma kitabını üretir ve `out_path`e kaydeder."""
    workouts = sorted(workouts or [], key=lambda w: (str(w["date"]), w.get("id") or 0))
    goals = goals or []
    settings = settings or {}
    athlete = settings.get("athlete_name") or "Sporcu"
    start_date = _to_date(settings.get("start_date")) or dt.date.today()
    if workouts:
        first = _to_date(workouts[0]["date"])
        if first:
            start_date = min(start_date, first)

    n_log = max(600, len(workouts) + 200)
    log_last = 1 + n_log
    n_brick = max(100, len({w["date"] for w in workouts if w["is_brick"]}) + 50)

    thr = thresholds or {}
    sw_t = thr.get("swim", {}).get("distance_km", [0.05, 0.5, 0.75, 1.5, 1.9, 3.0, 3.8])
    bk_t = thr.get("bike", {}).get("distance_km", [10, 30, 30, 40, 90, 120, 180])
    rn_t = thr.get("run", {}).get("distance_km", [1, 5, 5, 10, 21.1, 25, 42.2])
    sw_v = thr.get("swim", {}).get("volume_30d_km", [0.2, 2, 4, 8, 12, 20, 28])
    bk_v = thr.get("bike", {}).get("volume_30d_km", [20, 60, 100, 160, 280, 400, 500])
    rn_v = thr.get("run", {}).get("volume_30d_km", [4, 20, 30, 50, 80, 110, 140])

    wb = Workbook()
    ws_dash = wb.active
    ws_dash.title = SH_DASH
    ws_lvl = wb.create_sheet(SH_LVL)
    ws_log = wb.create_sheet(SH_LOG)
    ws_week = wb.create_sheet(SH_WEEK)
    ws_br = wb.create_sheet(SH_BR)
    ws_brick = wb.create_sheet(SH_BRICK)
    ws_goal = wb.create_sheet(SH_GOAL)
    ws_help = wb.create_sheet(SH_HELP)
    ws_set = wb.create_sheet(SH_SET)

    for ws in wb.worksheets:
        ws.sheet_properties.tabColor = NAVY
    ws_dash.sheet_properties.tabColor = IRON
    ws_log.sheet_properties.tabColor = SLATE
    ws_help.sheet_properties.tabColor = OK_G
    ws_set.sheet_properties.tabColor = MUTED

    _sheet_settings(ws_set, start_date, sw_t, bk_t, rn_t, sw_v, bk_v, rn_v)
    _defined_names(wb, log_last)
    _sheet_log(ws_log, workouts, log_last)
    _sheet_branch(ws_br)
    _sheet_levels(ws_lvl)
    _sheet_weekly(ws_week, n_week)
    _sheet_brick(ws_brick, workouts, n_brick)
    _sheet_goals(ws_goal, goals)
    _sheet_chartdata(ws_set)
    _sheet_dashboard(ws_dash, athlete, len(goals))
    _sheet_help(ws_help, athlete, sw_t, bk_t, rn_t)

    wb.active = 0
    title_rows = {SH_LOG: "1:1", SH_WEEK: "1:1", SH_LVL: "3:3",
                  SH_BRICK: "3:3", SH_GOAL: "3:3"}
    for ws in wb.worksheets:
        ws.page_setup.orientation = "landscape"
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
        ws.print_options.horizontalCentered = True
        if ws.title in title_rows:
            ws.print_title_rows = title_rows[ws.title]
    wb.calculation.fullCalcOnLoad = True

    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    _patch_full_calc(out_path)
    return out_path


def _patch_full_calc(path: Path) -> None:
    """openpyxl bazı sürümlerde calcPr'yi yazmaz; garanti altına al."""
    import re, shutil, zipfile
    tmp = path.with_suffix(".tmp.xlsx")
    with zipfile.ZipFile(path) as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "xl/workbook.xml":
                xml = data.decode("utf-8")
                if "<calcPr" in xml:
                    def rep(m):
                        tag = re.sub(r'\s*fullCalcOnLoad="[^"]*"', "", m.group(0))
                        return tag.rstrip("/>").rstrip() + ' fullCalcOnLoad="1"/>'
                    xml = re.sub(r"<calcPr[^>]*/?>", rep, xml, count=1)
                else:
                    xml = xml.replace("</workbook>",
                                      '<calcPr calcId="191029" fullCalcOnLoad="1"/></workbook>')
                data = xml.encode("utf-8")
            zout.writestr(item, data)
    shutil.move(str(tmp), str(path))


# ==========================================================================
# AYARLAR
# ==========================================================================
def _sheet_settings(S, start_date, sw_t, bk_t, rn_t, sw_v, bk_v, rn_v):
    S.sheet_view.showGridLines = False
    widths(S, {"A": 20, "B": 21, "C": 15, "D": 19, "E": 8, "F": 3, "G": 9, "H": 17,
               "I": 19, "J": 15, "K": 19, "L": 21, "M": 17, "N": 3, "O": 9, "P": 26})
    merge_set(S, 1, 1, 1, 16, "⚙  AYARLAR & REFERANS TABLOLARI",
              f(14, True, WHITE), al("left", "center", indent=1), NAVY)
    S.row_dimensions[1].height = 30

    S["A2"] = "Takip Başlangıç Tarihi"
    S["A2"].font = f(10, True)
    S["B2"] = start_date
    S["B2"].number_format = FMT_DATE
    S["B2"].font = f(10, True, INPUT_FONT)
    S["B2"].fill = fill(INPUT_FILL)
    S["B2"].border = BOX
    S["B2"].alignment = al("center")
    S["C2"] = "◀ HAFTALIK TAKİP sayfasının başlangıç haftasını belirler (Pazartesi'ye yuvarlanır)."
    S["C2"].font = f(9, False, MUTED)

    header_row(S, 3, ["Branş", "Antrenman Türü", "Evet / Hayır", "Hedef Durumu", "RPE"], 1, 26)
    for i, v in enumerate([L_SWIM, L_BIKE, L_RUN]):
        S.cell(row=4 + i, column=1, value=v)
    for i, v in enumerate(WORKOUT_KINDS):
        S.cell(row=4 + i, column=2, value=v)
    for i, v in enumerate(["Evet", "Hayır"]):
        S.cell(row=4 + i, column=3, value=v)
    for i, v in enumerate(["⬜ Başlanmadı", "🔄 Devam Ediyor", "✅ Tamamlandı"]):
        S.cell(row=4 + i, column=4, value=v)
    for i in range(10):
        S.cell(row=4 + i, column=5, value=i + 1)
    for r in range(4, 14):
        for c in range(1, 6):
            cell = S.cell(row=r, column=c)
            cell.font = f(10)
            cell.border = BOX
            cell.alignment = al("center" if c in (3, 5) else "left", "center",
                                indent=0 if c in (3, 5) else 1)
    S.cell(row=4, column=1).fill = fill(SWIM_L)
    S.cell(row=5, column=1).fill = fill(BIKE_L)
    S.cell(row=6, column=1).fill = fill(RUN_L)
    merge_set(S, 15, 1, 15, 5,
              "⚠️ A4:A6 branş etiketlerini DEĞİŞTİRMEYİN — dosyadaki tüm formüller bu üç etikete bağlıdır.",
              f(9, True, IRON), al("left", "center", True, indent=1), IRON_L)
    S.row_dimensions[15].height = 24

    header_row(S, 3, ["Seviye", "Yüzme Kriter (km)", "Bisiklet Kriter (km)", "Koşu Kriter (km)",
                      "Yüzme 30 Gün Hacim (km)", "Bisiklet 30 Gün Hacim (km)",
                      "Koşu 30 Gün Hacim (km)"], 7, 26)
    cols = [sw_t, bk_t, rn_t, sw_v, bk_v, rn_v]
    for i in range(7):
        S.cell(row=4 + i, column=7, value=i).number_format = FMT_INT
        for j, series in enumerate(cols):
            cell = S.cell(row=4 + i, column=8 + j, value=float(series[i]))
            cell.number_format = FMT_KM2 if j in (0, 3) else FMT_KM1
        for c in range(7, 14):
            cell = S.cell(row=4 + i, column=c)
            cell.font = f(10, c == 7, INPUT_FONT if c > 7 else INK)
            cell.border = BOX
            cell.alignment = al("center")
            if c > 7:
                cell.fill = fill(INPUT_FILL)
    S["G11"] = ("ℹ️ Sarı hücreler değiştirilebilir eşiklerdir. Seviye motoru bu tabloyu kullanır; "
                "kriter sütunları artan sırada kalmalıdır.")
    S["G11"].font = f(9, False, MUTED)
    S.merge_cells("G11:M12")
    S["G11"].alignment = al("left", "center", True)

    header_row(S, 3, ["Seviye", "Seviye Adı"], 15, 26)
    names = LEVEL_NAMES[:6] + ["IRONMAN 🏆"]
    for i, a in enumerate(names):
        S.cell(row=4 + i, column=15, value=i).number_format = FMT_INT
        S.cell(row=4 + i, column=16, value=a)
        for c in (15, 16):
            cell = S.cell(row=4 + i, column=c)
            cell.font = f(10, c == 16)
            cell.border = BOX
            cell.alignment = al("center" if c == 15 else "left", "center",
                                indent=0 if c == 15 else 1)


def _defined_names(wb, log_last):
    q = f"'{SH_LOG}'"
    names = {
        "Tarih": f"{q}!$A$2:$A${log_last}", "Brans": f"{q}!$B$2:$B${log_last}",
        "Tur": f"{q}!$C$2:$C${log_last}", "Mesafe": f"{q}!$D$2:$D${log_last}",
        "Sure": f"{q}!$E$2:$E${log_last}", "Nabiz": f"{q}!$H$2:$H${log_last}",
        "RPE": f"{q}!$I$2:$I${log_last}", "Yukselti": f"{q}!$J$2:$J${log_last}",
        "Kalori": f"{q}!$K$2:$K${log_last}", "Brick": f"{q}!$L$2:$L${log_last}",
        "AcikSu": f"{q}!$M$2:$M${log_last}", "HaftaBasi": f"{q}!$O$2:$O${log_last}",
        "B_Yuzme": f"{SH_SET}!$A$4", "B_Bisiklet": f"{SH_SET}!$A$5", "B_Kosu": f"{SH_SET}!$A$6",
    }
    for n, ref in names.items():
        wb.defined_names[n] = DefinedName(n, attr_text=ref)


# ==========================================================================
# ANTRENMAN KAYIT
# ==========================================================================
def _sheet_log(W, workouts, log_last):
    head = ["Tarih", "Branş", "Antrenman Türü", "Mesafe (km)", "Süre (dk)",
            "Pace  (🏊 dk/100m · 🏃 dk/km)", "Ort. Hız (km/s)", "Ort. Nabız",
            "RPE (1-10)", "Yükselti (m)", "Kalori", "Brick mi?", "Açık Su mu?",
            "Notlar", "⚙ Hafta Başı"]
    header_row(W, 1, head, 1, 42)
    widths(W, {"A": 12.5, "B": 14, "C": 15, "D": 11.5, "E": 10.5, "F": 14.5, "G": 12.5,
               "H": 11, "I": 11, "J": 12, "K": 10, "L": 10.5, "M": 12, "N": 42, "O": 12})
    W.freeze_panes = "D2"

    for r in range(2, log_last + 1):
        W.cell(row=r, column=6).value = (
            f'=IF(OR($D{r}="",$D{r}=0,$E{r}=""),"",'
            f'IF($B{r}=B_Yuzme,($E{r}/($D{r}*10))/1440,'
            f'IF($B{r}=B_Kosu,($E{r}/$D{r})/1440,"")))')
        W.cell(row=r, column=7).value = (
            f'=IF(OR($D{r}="",$D{r}=0,$E{r}="",$E{r}=0),"",$D{r}/($E{r}/60))')
        W.cell(row=r, column=15).value = f'=IF($A{r}="","",$A{r}-WEEKDAY($A{r},3))'

    fmts = {1: FMT_DATE, 4: FMT_KM2, 5: FMT_INT, 6: FMT_PACE, 7: FMT_SPD,
            8: FMT_INT, 9: FMT_INT, 10: FMT_INT, 11: FMT_INT, 15: FMT_DATE}
    for r in range(2, log_last + 1):
        for c in range(1, 16):
            cell = W.cell(row=r, column=c)
            cell.number_format = fmts.get(c, "General")
            cell.border = Border(bottom=side("EEF2F7"), right=side("EEF2F7"))
            cell.alignment = (al("left", "center", indent=1) if c in (2, 3, 12, 13, 14)
                              else al("center", "center"))
            if c in (1, 2, 3, 4, 5, 8, 9, 10, 11, 12, 13, 14):
                cell.font = f(10, False, INPUT_FONT)
                cell.fill = fill(INPUT_FILL)
            else:
                cell.font = f(10, False, MUTED)
                cell.fill = fill(GREY_BG)
        W.row_dimensions[r].height = 17

    # gerçek veriler
    for i, w in enumerate(workouts):
        r = 2 + i
        W.cell(row=r, column=1, value=_to_date(w["date"]))
        W.cell(row=r, column=2, value=SPORT_LABEL.get(w["sport"], w["sport"]))
        W.cell(row=r, column=3, value=w.get("kind") or "")
        W.cell(row=r, column=4, value=float(w["distance_km"] or 0))
        W.cell(row=r, column=5, value=float(w["duration_min"] or 0))
        W.cell(row=r, column=8, value=w.get("avg_hr"))
        W.cell(row=r, column=9, value=w.get("rpe"))
        W.cell(row=r, column=10, value=w.get("elevation_m"))
        W.cell(row=r, column=11, value=w.get("calories"))
        W.cell(row=r, column=12, value="Evet" if w.get("is_brick") else "Hayır")
        W.cell(row=r, column=13, value="Evet" if w.get("is_openwater") else "Hayır")
        W.cell(row=r, column=14, value=(w.get("notes") or "")[:400])

    W.column_dimensions["O"].hidden = True
    W.auto_filter.ref = f"A1:N{log_last}"

    dvs = [
        (DataValidation(type="list", formula1=f"{SH_SET}!$A$4:$A$6", allow_blank=True,
                        showDropDown=False), f"B2:B{log_last}"),
        (DataValidation(type="list", formula1=f"{SH_SET}!$B$4:$B$13", allow_blank=True,
                        showDropDown=False), f"C2:C{log_last}"),
        (DataValidation(type="list", formula1=f"{SH_SET}!$E$4:$E$13", allow_blank=True,
                        showDropDown=False), f"I2:I{log_last}"),
        (DataValidation(type="list", formula1=f"{SH_SET}!$C$4:$C$5", allow_blank=True,
                        showDropDown=False), f"L2:L{log_last}"),
        (DataValidation(type="list", formula1=f"{SH_SET}!$C$4:$C$5", allow_blank=True,
                        showDropDown=False), f"M2:M{log_last}"),
        (DataValidation(type="decimal", operator="greaterThanOrEqual", formula1="0",
                        allow_blank=True, errorTitle="Geçersiz mesafe",
                        error="Mesafeyi km cinsinden girin (750 m → 0,75)"), f"D2:D{log_last}"),
        (DataValidation(type="decimal", operator="greaterThanOrEqual", formula1="0",
                        allow_blank=True, errorTitle="Geçersiz süre",
                        error="Süreyi dakika cinsinden girin (1s 12dk → 72)"), f"E2:E{log_last}"),
    ]
    for dv, rng in dvs:
        dv.showErrorMessage = True
        if dv.type == "list" and not dv.errorTitle:
            dv.errorTitle = "Listeden seçin"
            dv.error = ("Bu hücreye sadece açılır listedeki değerler girilebilir. "
                        "Formüller bu etiketlere göre çalışır.")
        W.add_data_validation(dv)
        dv.add(rng)

    W["D1"].comment = Comment(
        "Mesafeyi HER BRANŞ İÇİN KİLOMETRE olarak girin.\n"
        "  • Yüzme 750 m  →  0,75\n  • Bisiklet 45 km →  45\n"
        "Pace ve hız otomatik hesaplanır.", "IRONMAN Takip", height=110, width=260)
    W["E1"].comment = Comment(
        "Süreyi DAKİKA olarak girin.\n  • 1 saat 12 dk  →  72\n  • 3 saat 30 dk  →  210",
        "IRONMAN Takip", height=90, width=230)


# ==========================================================================
# BRANŞ İLERLEME
# ==========================================================================
BR_ROWS = [
    (4, "Toplam Antrenman Sayısı", FMT_INT, "Bu branşta kaydedilen toplam antrenman"),
    (5, "Toplam Mesafe (km)", FMT_KM1, "Tüm zamanların toplam mesafesi"),
    (6, "Toplam Süre (saat)", FMT_KM1, "Tüm zamanların toplam antrenman süresi"),
    (7, "Mevcut Maksimum Mesafe (km)", FMT_KM2, "★ En uzun tek antrenman — seviye motorunun ana girdisi"),
    (8, "En Uzun Antrenmanın Tarihi", FMT_DATE, "Maksimum mesafeye ilk ulaşılan tarih"),
    (9, "Ortalama Mesafe / Antrenman (km)", FMT_KM2, "Toplam mesafe ÷ antrenman sayısı"),
    (10, "Ortalama Pace", FMT_PACE, "🏊 dk/100 m · 🏃 dk/km (bisiklette hız kullanılır)"),
    (11, "Ortalama Hız (km/s)", FMT_SPD, "Toplam mesafe ÷ toplam süre"),
    (12, "Son 7 Gün Mesafe (km)", FMT_KM2, "Bugün dahil son 7 gün"),
    (13, "Son 30 Gün Mesafe (km)", FMT_KM2, "★ Süreklilik/hacim seviyesinin girdisi"),
    (14, "Son 30 Gün Antrenman Sayısı", FMT_INT, "Bugün dahil son 30 gün"),
    (15, "Son 90 Gün Mesafe (km)", FMT_KM2, "Uzun vadeli hacim trendi"),
    (17, "Mesafe Seviyesi  (Rozet — düşmez)", FMT_INT, "En uzun antrenmana göre ulaşılan en yüksek seviye"),
    (18, "Form Seviyesi  (son 30 gün hacim)", FMT_INT, "Süreklilik göstergesi — antrenmana ara verilirse düşer"),
    (19, "MEVCUT SEVİYE", FMT_INT, "MIN(Mesafe Seviyesi ; Form Seviyesi) — ikisinin küçüğü"),
    (20, "Seviye Adı", "General", "Mevcut seviyenin adı"),
    (21, "Sonraki Seviye", "General", "Rozet seviyesinin bir üstü"),
    (22, "Sonraki Seviye Hedefi (km)", FMT_KM2, "Bir üst seviyeye geçmek için gereken mesafe"),
    (23, "Hedefe Kalan (km)", FMT_KM2, "Sonraki seviye hedefi − mevcut maksimum"),
    (24, "Seviye İçi İlerleme %", FMT_PCT, "Mevcut seviye ile sonraki seviye arasındaki ilerleme"),
    (25, "IRONMAN'E İLERLEME %  (seviye)", FMT_PCT, "★ (Seviye + seviye içi ilerleme) ÷ 6"),
    (26, "Ironman Yarış Mesafesi (km)", FMT_KM2, "3,8 km · 180 km · 42,2 km"),
    (27, "Ironman Mesafesine Oran %", FMT_PCT, "Maksimum mesafe ÷ Ironman mesafesi (ham oran)"),
]
BR_COLCFG = {"B": ("H", "K", SWIM, SWIM_XL), "C": ("I", "L", BIKE, BIKE_XL),
             "D": ("J", "M", RUN, RUN_XL)}


def _sheet_branch(B):
    B.sheet_view.showGridLines = False
    widths(B, {"A": 36, "B": 17, "C": 17, "D": 17, "E": 3, "F": 62})
    merge_set(B, 1, 1, 1, 6, "📈  BRANŞ İLERLEME — Detaylı Analiz",
              f(14, True, WHITE), al("left", "center", indent=1), NAVY)
    B.row_dimensions[1].height = 30
    merge_set(B, 2, 1, 2, 6, "Tüm değerler ANTRENMAN KAYIT sayfasından otomatik hesaplanır.",
              f(9, False, MUTED), al("left", "center", indent=1))
    header_row(B, 3, ["Metrik", "", "", "", "", "Açıklama"], 1, 30)
    for col, nm, c in ((2, "B_Yuzme", SWIM), (3, "B_Bisiklet", BIKE), (4, "B_Kosu", RUN)):
        cell = B.cell(row=3, column=col, value=f"={nm}")
        cell.font = f(11, True, WHITE)
        cell.fill = fill(c)
        cell.alignment = al("center")
    B.cell(row=3, column=5).fill = fill(WHITE)
    B.cell(row=3, column=5).border = Border()

    for r, label, numfmt, note in BR_ROWS:
        c = B.cell(row=r, column=1, value=label)
        c.font = f(10, r in (7, 13, 19, 25))
        c.alignment = al("left", "center", indent=1)
        c.fill = fill(WHITE if r % 2 else BG)
        c.border = Border(bottom=side("EEF2F7"))
        n = B.cell(row=r, column=6, value=note)
        n.font = f(9, False, MUTED)
        n.alignment = al("left", "center", indent=1)
        B.row_dimensions[r].height = 19

    for col, (T, V, acc, light) in BR_COLCFG.items():
        fml = {
            4: f'=COUNTIFS(Brans,{col}$3)',
            5: f'=SUMIFS(Mesafe,Brans,{col}$3)',
            6: f'=SUMIFS(Sure,Brans,{col}$3)/60',
            7: f'=SUMPRODUCT(MAX((Brans={col}$3)*Mesafe))',
            8: (f'=IF({col}$7=0,"",SUMPRODUCT(MIN((Brans={col}$3)*(Mesafe>={col}$7)*Tarih'
                f'+(1-(Brans={col}$3)*(Mesafe>={col}$7))*100000)))'),
            9: f'=IFERROR({col}$5/{col}$4,0)',
            11: f'=IF({col}$6=0,"",{col}$5/{col}$6)',
            12: f'=SUMIFS(Mesafe,Brans,{col}$3,Tarih,">="&TODAY()-6,Tarih,"<="&TODAY())',
            13: f'=SUMIFS(Mesafe,Brans,{col}$3,Tarih,">="&TODAY()-29,Tarih,"<="&TODAY())',
            14: f'=COUNTIFS(Brans,{col}$3,Tarih,">="&TODAY()-29,Tarih,"<="&TODAY())',
            15: f'=SUMIFS(Mesafe,Brans,{col}$3,Tarih,">="&TODAY()-89,Tarih,"<="&TODAY())',
            17: f'=MAX(0,SUMPRODUCT(--({col}$7>=AYARLAR!${T}$4:${T}$10))-1)',
            18: f'=MAX(0,SUMPRODUCT(--({col}$13>=AYARLAR!${V}$4:${V}$10))-1)',
            19: f'=MIN({col}$17,{col}$18)',
            20: f'=INDEX(AYARLAR!$P$4:$P$10,{col}$19+1)',
            21: (f'=IF({col}$17>=6,"🏆 Maksimum seviye","Seviye "&({col}$17+1)&" — "'
                 f'&INDEX(AYARLAR!$P$4:$P$10,{col}$17+2))'),
            22: f'=IF({col}$17>=6,AYARLAR!${T}$10,INDEX(AYARLAR!${T}$4:${T}$10,{col}$17+2))',
            23: f'=MAX(0,{col}$22-{col}$7)',
            24: (f'=IF({col}$17>=6,1,IFERROR(MIN(1,MAX(0,({col}$7-INDEX(AYARLAR!${T}$4:${T}$10,{col}$17+1))'
                 f'/({col}$22-INDEX(AYARLAR!${T}$4:${T}$10,{col}$17+1)))),1))'),
            25: f'=MIN(1,({col}$17+{col}$24)/6)',
            26: f'=AYARLAR!${T}$10',
            27: f'=MIN(1,{col}$7/{col}$26)',
        }
        fml[10] = ('=IF($B$5=0,"",($B$6*60/($B$5*10))/1440)' if col == "B" else
                   '=IF($D$5=0,"",($D$6*60/$D$5)/1440)' if col == "D" else '="—"')
        for r, label, numfmt, note in BR_ROWS:
            cell = B[f"{col}{r}"]
            cell.value = fml[r]
            cell.number_format = numfmt
            cell.alignment = al("center")
            cell.border = Border(bottom=side("EEF2F7"), left=side(LINE), right=side(LINE))
            if r in (19, 25):
                cell.font, cell.fill = f(12, True, acc), fill(light)
            elif r in (7, 13, 17):
                cell.font, cell.fill = f(11, True, INK), fill(light)
            elif r in (20, 21):
                cell.font = f(9.5, True, acc)
                cell.fill = fill(WHITE if r % 2 else BG)
                cell.alignment = al("center", "center", True)
            else:
                cell.font, cell.fill = f(10), fill(WHITE if r % 2 else BG)

    merge_set(B, 16, 1, 16, 6, "SEVİYE DURUMU", f(10, True, WHITE),
              al("left", "center", indent=1), SLATE)
    B.row_dimensions[16].height = 22
    B.row_dimensions[21].height = 30
    for col in ("B", "C", "D"):
        B.conditional_formatting.add(f"{col}24", DataBarRule(
            start_type="num", start_value=0, end_type="num", end_value=1,
            color=BR_COLCFG[col][2]))
        B.conditional_formatting.add(f"{col}25", DataBarRule(
            start_type="num", start_value=0, end_type="num", end_value=1, color=IRON))
        B.conditional_formatting.add(f"{col}27", DataBarRule(
            start_type="num", start_value=0, end_type="num", end_value=1, color=MUTED))
    B.freeze_panes = "B4"


# ==========================================================================
# SEVİYELER
# ==========================================================================
def _sheet_levels(V):
    widths(V, {"A": 14, "B": 8, "C": 21, "D": 25, "E": 11, "F": 15, "G": 52,
               "H": 12, "I": 12, "J": 15, "K": 15, "L": 44})
    merge_set(V, 1, 1, 1, 12, "🎯  SEVİYE SİSTEMİ — Seviye 0'dan IRONMAN'e",
              f(14, True, WHITE), al("left", "center", indent=1), NAVY)
    V.row_dimensions[1].height = 30
    merge_set(V, 2, 1, 2, 12,
              "«Kriter (km)» sütunu seviye motorunun kullandığı eşiktir (AYARLAR sayfasından gelir). "
              "«Tamamlandı» ve «Tamamlanma Tarihi» otomatik hesaplanır.",
              f(9, False, MUTED), al("left", "center", indent=1))
    header_row(V, 3, ["Branş", "Seviye", "Seviye Adı", "Hedef Mesafe", "Kriter (km)",
                      "Hedef Süre", "Gereken Kriter", "Mevcut (km)", "İlerleme %",
                      "Tamamlandı mı?", "Tamamlanma Tarihi", "Not"], 1, 34)

    thrcol = {"swim": "H", "bike": "I", "run": "J"}
    colors = {"swim": (SWIM, SWIM_XL), "bike": (BIKE, BIKE_XL), "run": (RUN, RUN_XL)}
    r = 4
    for sport in SPORT_ORDER:
        acc, light = colors[sport]
        T = thrcol[sport]
        label = SPORT_LABEL[sport]
        for lvl in range(7):
            hedef, sure, kriter, not_ = LEVEL_DETAIL[sport][lvl]
            V.cell(row=r, column=1, value=label)
            V.cell(row=r, column=2, value=lvl)
            V.cell(row=r, column=3, value=f"=INDEX(AYARLAR!$P$4:$P$10,{lvl}+1)")
            V.cell(row=r, column=4, value=hedef)
            V.cell(row=r, column=5, value=f"=AYARLAR!${T}${4 + lvl}")
            V.cell(row=r, column=6, value=sure)
            V.cell(row=r, column=7, value=kriter)
            V.cell(row=r, column=8, value=f'=SUMPRODUCT(MAX((Brans=$A{r})*Mesafe))')
            V.cell(row=r, column=9, value=f'=IFERROR(MIN(1,$H{r}/$E{r}),0)')
            V.cell(row=r, column=10, value=f'=IF($H{r}>=$E{r},"✅ Evet","⬜ Hayır")')
            V.cell(row=r, column=11, value=(
                f'=IF($H{r}<$E{r},"",SUMPRODUCT(MIN((Brans=$A{r})*(Mesafe>=$E{r})*Tarih'
                f'+(1-(Brans=$A{r})*(Mesafe>=$E{r}))*100000)))'))
            V.cell(row=r, column=12, value=not_)
            V.row_dimensions[r].height = 26
            for c in range(1, 13):
                cell = V.cell(row=r, column=c)
                cell.border = Border(bottom=side(LINE), left=side("EEF2F7"), right=side("EEF2F7"))
                cell.font = f(10)
                cell.alignment = (al("left", "center", True, indent=1) if c in (4, 7, 12)
                                  else al("center"))
                cell.fill = fill(light if lvl % 2 == 0 else WHITE)
            V.cell(row=r, column=1).font = f(10, True, acc)
            V.cell(row=r, column=2).font = f(12, True, acc)
            V.cell(row=r, column=3).font = f(10, True)
            V.cell(row=r, column=5).number_format = FMT_KM2
            V.cell(row=r, column=8).number_format = FMT_KM2
            V.cell(row=r, column=9).number_format = FMT_PCT
            V.cell(row=r, column=11).number_format = FMT_DATE
            V.cell(row=r, column=12).font = f(9, False, MUTED)
            V.cell(row=r, column=7).font = f(9.5)
            r += 1

    last = r - 1
    V.freeze_panes = "D4"
    V.auto_filter.ref = f"A3:L{last}"
    V.conditional_formatting.add(f"I4:I{last}", DataBarRule(
        start_type="num", start_value=0, end_type="num", end_value=1, color=OK_G))
    V.conditional_formatting.add(f"A4:L{last}", FormulaRule(
        formula=["$H4>=$E4"], fill=fill(OK_BG), font=f(10, False, "14532D")))
    V.conditional_formatting.add(f"J4:J{last}", FormulaRule(
        formula=["$H4>=$E4"], fill=fill(OK_G), font=f(10, True, WHITE)))


# ==========================================================================
# HAFTALIK TAKİP
# ==========================================================================
def _sheet_weekly(H, n_week):
    head = ["Hafta Başlangıcı", "Hafta", "🏊 Yüzme (km)", "🚴 Bisiklet (km)", "🏃 Koşu (km)",
            "Toplam Mesafe (km)", "Toplam Süre (saat)", "Antrenman Sayısı",
            "Uzun Koşu (km)", "Uzun Bisiklet (km)", "En Uzun Yüzme (km)", "Brick Sayısı",
            "Süre Δ % (önceki hafta)", "Mesafe Δ % (önceki hafta)", "Not"]
    header_row(H, 1, head, 1, 44)
    widths(H, {"A": 16, "B": 8, "C": 13, "D": 14, "E": 12, "F": 15, "G": 15, "H": 13,
               "I": 13, "J": 15, "K": 15, "L": 12, "M": 15, "N": 15, "O": 40})
    H.freeze_panes = "C2"
    last = 1 + n_week
    for i in range(n_week):
        r = 2 + i
        H.cell(row=r, column=1, value=("=AYARLAR!$B$2-WEEKDAY(AYARLAR!$B$2,3)" if i == 0
                                       else f"=A{r-1}+7"))
        H.cell(row=r, column=2, value='="H"&(ROW()-1)')
        H.cell(row=r, column=3, value=f'=IF($A{r}>TODAY(),"",SUMIFS(Mesafe,HaftaBasi,$A{r},Brans,B_Yuzme))')
        H.cell(row=r, column=4, value=f'=IF($A{r}>TODAY(),"",SUMIFS(Mesafe,HaftaBasi,$A{r},Brans,B_Bisiklet))')
        H.cell(row=r, column=5, value=f'=IF($A{r}>TODAY(),"",SUMIFS(Mesafe,HaftaBasi,$A{r},Brans,B_Kosu))')
        H.cell(row=r, column=6, value=f'=IF($A{r}>TODAY(),"",SUM($C{r}:$E{r}))')
        H.cell(row=r, column=7, value=f'=IF($A{r}>TODAY(),"",SUMIFS(Sure,HaftaBasi,$A{r})/60)')
        H.cell(row=r, column=8, value=f'=IF($A{r}>TODAY(),"",COUNTIFS(HaftaBasi,$A{r}))')
        H.cell(row=r, column=9, value=f'=IF($A{r}>TODAY(),"",SUMPRODUCT(MAX((HaftaBasi=$A{r})*(Brans=B_Kosu)*Mesafe)))')
        H.cell(row=r, column=10, value=f'=IF($A{r}>TODAY(),"",SUMPRODUCT(MAX((HaftaBasi=$A{r})*(Brans=B_Bisiklet)*Mesafe)))')
        H.cell(row=r, column=11, value=f'=IF($A{r}>TODAY(),"",SUMPRODUCT(MAX((HaftaBasi=$A{r})*(Brans=B_Yuzme)*Mesafe)))')
        H.cell(row=r, column=12, value=f'=IF($A{r}>TODAY(),"",COUNTIFS(HaftaBasi,$A{r},Brick,"Evet",Brans,B_Bisiklet))')
        if i == 0:
            H.cell(row=r, column=13, value='=""')
            H.cell(row=r, column=14, value='=""')
        else:
            H.cell(row=r, column=13, value=f'=IF(OR($A{r}>TODAY(),N($G{r-1})=0),"",$G{r}/$G{r-1}-1)')
            H.cell(row=r, column=14, value=f'=IF(OR($A{r}>TODAY(),N($F{r-1})=0),"",$F{r}/$F{r-1}-1)')
        H.row_dimensions[r].height = 17
        for c in range(1, 16):
            cell = H.cell(row=r, column=c)
            cell.border = Border(bottom=side("EEF2F7"), right=side("EEF2F7"))
            cell.alignment = al("center") if c != 15 else al("left", "center", indent=1)
            cell.font = f(10)
            cell.fill = fill(WHITE if i % 2 == 0 else BG)
        H.cell(row=r, column=1).number_format = FMT_DATE
        H.cell(row=r, column=1).font = f(10, True)
        H.cell(row=r, column=2).font = f(10, True, MUTED)
        for c in (3, 4, 5, 6, 9, 10, 11):
            H.cell(row=r, column=c).number_format = FMT_KM2 if c in (3, 11) else FMT_KM1
        H.cell(row=r, column=7).number_format = FMT_KM1
        H.cell(row=r, column=8).number_format = FMT_INT
        H.cell(row=r, column=12).number_format = FMT_INT
        H.cell(row=r, column=13).number_format = FMT_DPCT
        H.cell(row=r, column=14).number_format = FMT_DPCT
        H.cell(row=r, column=15).font = f(10, False, INPUT_FONT)
        H.cell(row=r, column=15).fill = fill(INPUT_FILL)

    H.auto_filter.ref = f"A1:O{last}"
    for rng, col in ((f"C2:C{last}", SWIM), (f"D2:D{last}", BIKE),
                     (f"E2:E{last}", RUN), (f"G2:G{last}", SLATE)):
        H.conditional_formatting.add(rng, DataBarRule(start_type="min", end_type="max", color=col))
    H.conditional_formatting.add(f"M2:N{last}", CellIsRule(
        operator="greaterThan", formula=["0"], font=f(10, True, OK_G)))
    H.conditional_formatting.add(f"M2:N{last}", CellIsRule(
        operator="lessThan", formula=["0"], font=f(10, True, IRON)))
    H.conditional_formatting.add(f"A2:B{last}", FormulaRule(
        formula=["$A2>TODAY()"], font=f(10, False, "CBD5E1")))


# ==========================================================================
# BRICK
# ==========================================================================
def _sheet_brick(K, workouts, n_brick):
    merge_set(K, 1, 1, 1, 12, "🔁  BRICK ANTRENMANLARI  (bisiklet → koşu)",
              f(14, True, WHITE), al("left", "center", indent=1), NAVY)
    K.row_dimensions[1].height = 30
    merge_set(K, 2, 1, 2, 12,
              "Sadece TARİH ve NOT sütunlarını doldurun. Diğer her şey, ANTRENMAN KAYIT sayfasında "
              "«Brick mi? = Evet» işaretlenmiş aynı tarihli kayıtlardan otomatik gelir.",
              f(9, False, MUTED), al("left", "center", True, indent=1))
    K.row_dimensions[2].height = 26
    head = ["Tarih", "🚴 Bisiklet Mesafe (km)", "🚴 Bisiklet Süre (dk)", "🏃 Koşu Mesafe (km)",
            "🏃 Koşu Süre (dk)", "Bisiklet Sonrası Koşu Pace", "Normal Koşu Pace (60 gün ort.)",
            "Pace Farkı (sn/km)", "🚴 Ort. Hız (km/s)", "Toplam Süre (dk)", "RPE (ort.)", "Not"]
    header_row(K, 3, head, 1, 42)
    widths(K, {"A": 13, "B": 16, "C": 15, "D": 15, "E": 14, "F": 18, "G": 20,
               "H": 15, "I": 15, "J": 14, "K": 11, "L": 46})
    K.freeze_panes = "B4"

    last = 3 + n_brick
    for i in range(n_brick):
        r = 4 + i
        K.cell(row=r, column=2, value=f'=IF($A{r}="","",SUMIFS(Mesafe,Tarih,$A{r},Brans,B_Bisiklet,Brick,"Evet"))')
        K.cell(row=r, column=3, value=f'=IF($A{r}="","",SUMIFS(Sure,Tarih,$A{r},Brans,B_Bisiklet,Brick,"Evet"))')
        K.cell(row=r, column=4, value=f'=IF($A{r}="","",SUMIFS(Mesafe,Tarih,$A{r},Brans,B_Kosu,Brick,"Evet"))')
        K.cell(row=r, column=5, value=f'=IF($A{r}="","",SUMIFS(Sure,Tarih,$A{r},Brans,B_Kosu,Brick,"Evet"))')
        K.cell(row=r, column=6, value=f'=IF(OR($A{r}="",N($D{r})=0),"",($E{r}/$D{r})/1440)')
        K.cell(row=r, column=7, value=(
            f'=IF($A{r}="","",IFERROR((SUMIFS(Sure,Brans,B_Kosu,Brick,"<>Evet",Tarih,">="&$A{r}-59,Tarih,"<="&$A{r})'
            f'/SUMIFS(Mesafe,Brans,B_Kosu,Brick,"<>Evet",Tarih,">="&$A{r}-59,Tarih,"<="&$A{r}))/1440,""))'))
        K.cell(row=r, column=8, value=f'=IF(OR($F{r}="",$G{r}=""),"",($F{r}-$G{r})*86400)')
        K.cell(row=r, column=9, value=f'=IF(OR($A{r}="",N($C{r})=0),"",$B{r}/($C{r}/60))')
        K.cell(row=r, column=10, value=f'=IF($A{r}="","",$C{r}+$E{r})')
        K.cell(row=r, column=11, value=f'=IF($A{r}="","",IFERROR(AVERAGEIFS(RPE,Tarih,$A{r},Brick,"Evet"),""))')
        K.row_dimensions[r].height = 17
        for c in range(1, 13):
            cell = K.cell(row=r, column=c)
            cell.border = Border(bottom=side("EEF2F7"), right=side("EEF2F7"))
            cell.alignment = al("center") if c != 12 else al("left", "center", indent=1)
            cell.font = f(10)
            cell.fill = fill(WHITE if i % 2 == 0 else BG)
        K.cell(row=r, column=1).number_format = FMT_DATE
        K.cell(row=r, column=1).font = f(10, True, INPUT_FONT)
        K.cell(row=r, column=1).fill = fill(INPUT_FILL)
        K.cell(row=r, column=12).font = f(10, False, INPUT_FONT)
        K.cell(row=r, column=12).fill = fill(INPUT_FILL)
        for c, fmt in ((2, FMT_KM1), (3, FMT_INT), (4, FMT_KM1), (5, FMT_INT),
                       (6, FMT_PACE), (7, FMT_PACE), (9, FMT_SPD), (10, FMT_INT), (11, FMT_SPD)):
            K.cell(row=r, column=c).number_format = fmt
        K.cell(row=r, column=8).number_format = '+0" sn";-0" sn";0" sn"'

    days = sorted({str(w["date"])[:10] for w in workouts if w.get("is_brick")}, reverse=True)
    for i, day in enumerate(days[:n_brick]):
        K.cell(row=4 + i, column=1, value=_to_date(day))

    K.auto_filter.ref = f"A3:L{last}"
    K.conditional_formatting.add(f"H4:H{last}", CellIsRule(
        operator="greaterThan", formula=["0"], font=f(10, True, IRON)))
    K.conditional_formatting.add(f"H4:H{last}", CellIsRule(
        operator="lessThanOrEqual", formula=["0"], font=f(10, True, OK_G)))
    dv = DataValidation(type="date", operator="greaterThan", formula1="DATE(2000,1,1)",
                        allow_blank=True)
    K.add_data_validation(dv)
    dv.add(f"A4:A{last}")


# ==========================================================================
# HEDEFLER
# ==========================================================================
def _sheet_goals(G, goals):
    merge_set(G, 1, 1, 1, 10, "🎯  HEDEFLER — Kısa / Orta / Uzun Vade",
              f(14, True, WHITE), al("left", "center", indent=1), NAVY)
    G.row_dimensions[1].height = 30
    merge_set(G, 2, 1, 2, 10,
              "Sarı hücreler (Hedef Tarih ve triatlon yarışlarının tamamlanma tarihi) size aittir; "
              "diğer tüm sütunlar otomatik hesaplanır.",
              f(9, False, MUTED), al("left", "center", indent=1))
    header_row(G, 3, ["Hedef", "Branş", "Hedef Tarih", "Mevcut Durum", "Hedef Değer",
                      "İlerleme %", "Durum", "Tamamlanma Tarihi", "Kriter / Not", "⚙"], 1, 34)
    widths(G, {"A": 40, "B": 14, "C": 14, "D": 14, "E": 14, "F": 13, "G": 17,
               "H": 17, "I": 52, "J": 5})

    TRI = "🏅 Triatlon"
    colors = {"swim": (SWIM, SWIM_XL), "bike": (BIKE, BIKE_XL),
              "run": (RUN, RUN_XL), "tri": (IRON, IRON_L)}
    maxref = {"swim": "$B$7", "bike": "$C$7", "run": "$D$7"}
    rngname = {"swim": "B_Yuzme", "bike": "B_Bisiklet", "run": "B_Kosu"}

    first = 4
    if not goals:
        goals = []
    for i, g in enumerate(goals):
        r = first + i
        sport = g["sport"]
        acc, light = colors.get(sport, colors["tri"])
        label = SPORT_LABEL.get(sport, TRI)
        G.cell(row=r, column=1, value=g["name"])
        G.cell(row=r, column=2, value=label)
        G.cell(row=r, column=3, value=_to_date(g.get("target_date")))
        if g["goal_type"] == "distance":
            G.cell(row=r, column=4, value=f"={BRQ}!{maxref[sport]}")
            G.cell(row=r, column=5, value=float(g["target_value"] or 0))
            G.cell(row=r, column=8, value=(
                f'=IF($F{r}<1,"",SUMPRODUCT(MIN((Brans={rngname[sport]})*(Mesafe>=$E{r})*Tarih'
                f'+(1-(Brans={rngname[sport]})*(Mesafe>=$E{r}))*100000)))'))
            G.cell(row=r, column=4).number_format = FMT_KM2
            G.cell(row=r, column=5).number_format = FMT_KM2
            G.cell(row=r, column=8).number_format = FMT_DATE
            G.cell(row=r, column=7, value=(
                f'=IF($F{r}>=1,"✅ Tamamlandı",IF($F{r}>0,"🔄 Devam Ediyor","⬜ Başlanmadı"))'))
            G.cell(row=r, column=10, value=f'=IF($F{r}>=1,9999,ROW())')
        else:
            s = float(g.get("target_swim") or 0)
            b = float(g.get("target_bike") or 0)
            k = float(g.get("target_run") or 0)
            G.cell(row=r, column=4, value=(
                f'=(({BRQ}!$B$7>={s})+({BRQ}!$C$7>={b})+({BRQ}!$D$7>={k}))*1'))
            G.cell(row=r, column=5, value=3)
            G.cell(row=r, column=4).number_format = '0" / 3"'
            G.cell(row=r, column=5).number_format = '0" ayak"'
            G.cell(row=r, column=8, value=_to_date(g.get("completed_date")))
            G.cell(row=r, column=8).number_format = FMT_DATE
            G.cell(row=r, column=8).font = f(10, True, INPUT_FONT)
            G.cell(row=r, column=8).fill = fill(INPUT_FILL)
            G.cell(row=r, column=7, value=(
                f'=IF($H{r}<>"","✅ Tamamlandı",IF($F{r}>=1,"🟢 Yarışa Hazır",'
                f'IF($F{r}>0,"🔄 Devam Ediyor","⬜ Başlanmadı")))'))
            G.cell(row=r, column=10, value=f'=IF($H{r}<>"",9999,ROW())')
        G.cell(row=r, column=6, value=f'=IFERROR(MIN(1,$D{r}/$E{r}),0)')
        G.cell(row=r, column=9, value=g.get("note") or "")

        G.row_dimensions[r].height = 24
        for c in range(1, 11):
            cell = G.cell(row=r, column=c)
            cell.border = Border(bottom=side(LINE), left=side("EEF2F7"), right=side("EEF2F7"))
            cell.font = f(10)
            cell.alignment = al("center") if c not in (1, 9) else al("left", "center", indent=1)
            cell.fill = fill(WHITE if i % 2 == 0 else BG)
        G.cell(row=r, column=1).font = f(10.5, True)
        G.cell(row=r, column=2).font = f(10, True, acc)
        G.cell(row=r, column=2).fill = fill(light)
        G.cell(row=r, column=3).number_format = FMT_DATE
        G.cell(row=r, column=3).font = f(10, True, INPUT_FONT)
        G.cell(row=r, column=3).fill = fill(INPUT_FILL)
        G.cell(row=r, column=6).number_format = FMT_PCT
        G.cell(row=r, column=7).font = f(10, True)
        G.cell(row=r, column=9).font = f(9, False, MUTED)

    last = first + max(0, len(goals) - 1)
    G.freeze_panes = "B4"
    G.column_dimensions["J"].hidden = True
    if goals:
        G.auto_filter.ref = f"A3:I{last}"
        G.conditional_formatting.add(f"F{first}:F{last}", DataBarRule(
            start_type="num", start_value=0, end_type="num", end_value=1, color=IRON))
        G.conditional_formatting.add(f"A{first}:I{last}", FormulaRule(
            formula=[f"$J{first}=9999"], fill=fill(OK_BG), font=f(10, False, "14532D")))
        G.conditional_formatting.add(f"G{first}:G{last}", FormulaRule(
            formula=[f"$J{first}=9999"], fill=fill(OK_G), font=f(10, True, WHITE),
            stopIfTrue=True))
        G.conditional_formatting.add(f"G{first}:G{last}", FormulaRule(
            formula=[f"$F{first}>=1"], fill=fill("BBF7D0"), font=f(10, True, "14532D")))
        G.conditional_formatting.add(f"C{first}:C{last}", FormulaRule(
            formula=[f"AND($J{first}<>9999,$C{first}<TODAY())"],
            fill=fill(IRON_L), font=f(10, True, IRON)))


# ==========================================================================
# Grafik verisi (AYARLAR sayfasının alt bölümü)
# ==========================================================================
def _sheet_chartdata(S):
    merge_set(S, 17, 1, 17, 7, "📊  GRAFİK VERİSİ — SON 12 HAFTA  (otomatik, silmeyin)",
              f(10, True, WHITE), al("left", "center", indent=1), SLATE)
    S.row_dimensions[17].height = 22
    header_row(S, 18, ["Hafta Başı", "Etiket", "Toplam Süre (saat)", "🏊 Yüzme (km)",
                       "🚴 Bisiklet (km)", "🏃 Koşu (km)", "Toplam (km)"], 1, 30)
    for i in range(12):
        r = 19 + i
        S.cell(row=r, column=1, value="=TODAY()-WEEKDAY(TODAY(),3)-7*(30-ROW())").number_format = FMT_DATE
        S.cell(row=r, column=2, value=f'=DAY($A{r})&"."&MONTH($A{r})')
        S.cell(row=r, column=3, value=f"=SUMIFS(Sure,HaftaBasi,$A{r})/60").number_format = FMT_KM1
        S.cell(row=r, column=4, value=f"=SUMIFS(Mesafe,HaftaBasi,$A{r},Brans,B_Yuzme)").number_format = FMT_KM2
        S.cell(row=r, column=5, value=f"=SUMIFS(Mesafe,HaftaBasi,$A{r},Brans,B_Bisiklet)").number_format = FMT_KM1
        S.cell(row=r, column=6, value=f"=SUMIFS(Mesafe,HaftaBasi,$A{r},Brans,B_Kosu)").number_format = FMT_KM1
        S.cell(row=r, column=7, value=f"=SUM($D{r}:$F{r})").number_format = FMT_KM1
        for c in range(1, 8):
            cell = S.cell(row=r, column=c)
            cell.font, cell.border, cell.alignment = f(10), BOX, al("center")
            cell.fill = fill(WHITE if i % 2 == 0 else BG)

    merge_set(S, 32, 1, 32, 3, "📊  BRANŞ DAĞILIMI  (otomatik)", f(10, True, WHITE),
              al("left", "center", indent=1), SLATE)
    S.row_dimensions[32].height = 22
    header_row(S, 33, ["Branş", "Toplam Süre (saat)", "Toplam Mesafe (km)"], 1, 30)
    for i in range(3):
        r = 34 + i
        S.cell(row=r, column=1, value=f"=$A${4+i}")
        S.cell(row=r, column=2, value=f"=SUMIFS(Sure,Brans,$A{r})/60").number_format = FMT_KM1
        S.cell(row=r, column=3, value=f"=SUMIFS(Mesafe,Brans,$A{r})").number_format = FMT_KM1
        for c in range(1, 4):
            cell = S.cell(row=r, column=c)
            cell.font, cell.border, cell.alignment = f(10), BOX, al("center")

    merge_set(S, 38, 1, 38, 3, "📊  BRANŞ BAZLI IRONMAN İLERLEMESİ  (otomatik)",
              f(10, True, WHITE), al("left", "center", indent=1), SLATE)
    S.row_dimensions[38].height = 22
    header_row(S, 39, ["Branş", "Ironman'e İlerleme %", "Mesafe Oranı %"], 1, 30)
    for i, col in enumerate(("B", "C", "D")):
        r = 40 + i
        S.cell(row=r, column=1, value=f"=$A${4+i}")
        S.cell(row=r, column=2, value=f"={BRQ}!${col}$25").number_format = FMT_PCT
        S.cell(row=r, column=3, value=f"={BRQ}!${col}$27").number_format = FMT_PCT
        for c in range(1, 4):
            cell = S.cell(row=r, column=c)
            cell.font, cell.border, cell.alignment = f(10), BOX, al("center")


# ==========================================================================
# DASHBOARD
# ==========================================================================
def _sheet_dashboard(D, athlete, n_goals):
    D.sheet_view.showGridLines = False
    D.sheet_view.zoomScale = 90
    widths(D, {"A": 2.2, "B": 13, "C": 13, "D": 13, "E": 2.2, "F": 13, "G": 13, "H": 13,
               "I": 2.2, "J": 13, "K": 13, "L": 13, "M": 2.2, "N": 2.2})
    fill_range(D, 1, 1, 70, 14, BG)
    heights = {1: 6, 2: 40, 3: 22, 4: 10, 5: 24, 6: 16, 7: 32, 8: 17, 9: 10, 10: 24,
               11: 17, 12: 30, 13: 17, 14: 15, 15: 15, 16: 13, 17: 19, 18: 10, 19: 24,
               20: 16, 21: 28, 22: 17, 23: 7, 24: 16, 25: 28, 26: 17, 27: 7, 28: 16,
               29: 28, 30: 17, 31: 10, 32: 24, 33: 8}
    for r, h in heights.items():
        D.row_dimensions[r].height = h

    merge_set(D, 2, 2, 2, 12, "🏊  🚴  🏃      IRONMAN YOLCULUĞU",
              f(20, True, WHITE), al("left", "center", indent=1), NAVY)
    merge_set(D, 3, 2, 3, 12,
              f"{athlete} · Sıfırdan Full Ironman'e  ·  3,8 km yüzme + 180 km bisiklet + 42,2 km koşu",
              f(9.5, False, "94A3B8"), al("left", "center", indent=1), NAVY)

    def section(row, text):
        merge_set(D, row, 2, row, 12, text, f(10.5, True, SLATE), al("left", "bottom"))
        for c in range(2, 13):
            D.cell(row=row, column=c).border = Border(bottom=side(SLATE, "medium"))

    def card(row, c1, title, value, sub, accent, val_fmt=None, val_size=17, rows=3):
        c3 = c1 + 2
        fill_range(D, row, c1, row + rows - 1, c3, WHITE)
        merge_set(D, row, c1, row, c3, title, f(8, True, MUTED), al("left", "center", indent=1))
        v = merge_set(D, row + 1, c1, row + 1, c3, value, f(val_size, True, accent),
                      al("left", "center", indent=1))
        if val_fmt:
            v.number_format = val_fmt
        merge_set(D, row + 2, c1, row + 2, c3, sub, f(9, False, MUTED),
                  al("left", "center", indent=1))
        apply_box(D, row, c1, row + rows - 1, c3, color=LINE, left_accent=accent)
        return v

    section(5, "GENEL DURUM")
    lvl_form = f"MIN({BRQ}!$B$19,{BRQ}!$C$19,{BRQ}!$D$19)"
    lvl_badge = f"MIN({BRQ}!$B$17,{BRQ}!$C$17,{BRQ}!$D$17)"
    card(6, 2, "GENEL SEVİYE  (güncel form)", f'="SEVİYE "&{lvl_form}',
         f'=INDEX(AYARLAR!$P$4:$P$10,{lvl_form}+1)&"  ·  en zayıf branş belirler"',
         IRON, val_size=19)
    card(6, 6, "ROZET SEVİYESİ  (ulaşılan en yüksek)", f'="SEVİYE "&{lvl_badge}',
         f'=INDEX(AYARLAR!$P$4:$P$10,{lvl_badge}+1)&"  ·  bu seviye asla düşmez"',
         SLATE, val_size=19)
    card(6, 10, "IRONMAN'E TAHMİNİ İLERLEME",
         f"=AVERAGE({BRQ}!$B$25,{BRQ}!$C$25,{BRQ}!$D$25)",
         f'="Ham mesafe oranı: %"&ROUND(AVERAGE({BRQ}!$B$27,{BRQ}!$C$27,{BRQ}!$D$27)*100,1)',
         IRON, val_fmt=FMT_PCT, val_size=19)

    section(10, "BRANŞ SEVİYELERİ")
    for c1, title, col, acc, is_swim in ((2, "🏊  YÜZME", "B", SWIM, True),
                                         (6, "🚴  BİSİKLET", "C", BIKE, False),
                                         (10, "🏃  KOŞU", "D", RUN, False)):
        c3 = c1 + 2
        fill_range(D, 11, c1, 17, c3, WHITE)
        merge_set(D, 11, c1, 11, c3, title, f(10, True, WHITE), al("left", "center", indent=1), acc)
        merge_set(D, 12, c1, 12, c3, f'="SEVİYE "&{BRQ}!${col}$19', f(22, True, acc),
                  al("left", "center", indent=1))
        merge_set(D, 13, c1, 13, c3, f"={BRQ}!${col}$20", f(10.5, True, SLATE),
                  al("left", "center", indent=1))
        if is_swim:
            mx = f'="Maksimum: "&ROUND({BRQ}!${col}$7*1000,0)&" m"'
            nxt = (f'=IF({BRQ}!${col}$17>=6,"🏆 Tüm seviyeler tamamlandı","Sonraki: "'
                   f'&ROUND({BRQ}!${col}$22*1000,0)&" m  (kalan "&ROUND({BRQ}!${col}$23*1000,0)&" m)")')
        else:
            mx = f'="Maksimum: "&ROUND({BRQ}!${col}$7,1)&" km"'
            nxt = (f'=IF({BRQ}!${col}$17>=6,"🏆 Tüm seviyeler tamamlandı","Sonraki: "'
                   f'&ROUND({BRQ}!${col}$22,1)&" km  (kalan "&ROUND({BRQ}!${col}$23,1)&" km)")')
        merge_set(D, 14, c1, 14, c3, mx, f(9, False, INK), al("left", "center", indent=1))
        merge_set(D, 15, c1, 15, c3, nxt, f(9, False, MUTED), al("left", "center", indent=1))
        merge_set(D, 16, c1, 16, c3, "IRONMAN'E İLERLEME", f(7.5, True, MUTED),
                  al("left", "center", indent=1))
        merge_set(D, 17, c1, 17, c3, f"={BRQ}!${col}$25", f(10, True, acc),
                  al("right", "center", indent=1), numfmt=FMT_PCT)
        D.conditional_formatting.add(f"{get_column_letter(c1)}17", DataBarRule(
            start_type="num", start_value=0, end_type="num", end_value=1, color=acc))
        apply_box(D, 11, c1, 17, c3, color=LINE, left_accent=acc)

    section(19, "TOPLAM İSTATİSTİKLER")
    mon = "TODAY()-WEEKDAY(TODAY(),3)"
    p30 = 'Tarih,">="&TODAY()-29,Tarih,"<="&TODAY()'
    pp30 = 'Tarih,">="&TODAY()-59,Tarih,"<="&TODAY()-30'
    card(20, 2, "TOPLAM ANTRENMAN SÜRESİ", "=SUM(Sure)/60",
         '=COUNT(Tarih)&" antrenman  ·  "&ROUND(SUM(Mesafe),0)&" km toplam mesafe"',
         SLATE, val_fmt=FMT_HR)
    card(20, 6, "🏊  TOPLAM YÜZME", f"={BRQ}!$B$5",
         f'={BRQ}!$B$4&" antrenman  ·  maks "&ROUND({BRQ}!$B$7*1000,0)&" m"',
         SWIM, val_fmt=FMT_KMU2)
    card(20, 10, "🚴  TOPLAM BİSİKLET", f"={BRQ}!$C$5",
         f'={BRQ}!$C$4&" antrenman  ·  maks "&ROUND({BRQ}!$C$7,1)&" km"',
         BIKE, val_fmt=FMT_KMU1)
    card(24, 2, "🏃  TOPLAM KOŞU", f"={BRQ}!$D$5",
         f'={BRQ}!$D$4&" antrenman  ·  maks "&ROUND({BRQ}!$D$7,1)&" km"',
         RUN, val_fmt=FMT_KMU1)
    card(24, 6, "BU HAFTA", f"=COUNTIFS(HaftaBasi,{mon})",
         f'=ROUND(SUMIFS(Sure,HaftaBasi,{mon})/60,1)&" saat  ·  "'
         f'&ROUND(SUMIFS(Mesafe,HaftaBasi,{mon}),1)&" km"', SLATE, val_fmt=FMT_ADET)
    card(24, 10, "SON 30 GÜN", f"=SUMIFS(Sure,{p30})/60",
         f'="🏊 "&ROUND({BRQ}!$B$13,2)&"  ·  🚴 "&ROUND({BRQ}!$C$13,1)'
         f'&"  ·  🏃 "&ROUND({BRQ}!$D$13,1)&" km"', SLATE, val_fmt=FMT_HR)
    card(28, 2, "SON 30 GÜN DEĞİŞİM",
         f'=IF(SUMIFS(Sure,{pp30})=0,"—",SUMIFS(Sure,{p30})/SUMIFS(Sure,{pp30})-1)',
         f'=IF(SUMIFS(Sure,{pp30})=0,"önceki 30 günde kayıt yok",'
         f'"önceki 30 güne göre antrenman süresi")', OK_G, val_fmt=FMT_DPCT)
    glast = 3 + max(1, n_goals)
    nxt_i = f"MIN(HEDEFLER!$J$4:$J${glast})-3"
    card(28, 6, "SONRAKİ HEDEF",
         f'=IFERROR(INDEX(HEDEFLER!$A$4:$A${glast},{nxt_i}),"🏆 Tüm hedefler tamamlandı!")',
         f'=IFERROR("İlerleme %"&ROUND(INDEX(HEDEFLER!$F$4:$F${glast},{nxt_i})*100,0)'
         f'&"  ·  Hedef: "&{datestr(f"INDEX(HEDEFLER!$C$4:$C${glast},{nxt_i})")},"")',
         IRON, val_size=11)
    card(28, 10, "IRONMAN YARIŞ HEDEFİ", '="3,8 + 180 + 42,2 km"',
         f'="Mesafe bazlı hazırlık: %"&ROUND(AVERAGE({BRQ}!$B$27,{BRQ}!$C$27,{BRQ}!$D$27)*100,0)',
         IRON, val_size=14)
    D.conditional_formatting.add("B29", DataBarRule(
        start_type="num", start_value=-1, end_type="num", end_value=1, color=OK_G))

    section(32, "GRAFİKLER")
    S = D.parent[SH_SET]

    def dlabels(show_val=False, show_pct=False, numfmt=None, pos=None):
        d = DataLabelList()
        d.showVal, d.showPercent = show_val, show_pct
        d.showCatName = d.showSerName = d.showLegendKey = d.showBubbleSize = False
        if numfmt:
            d.numFmt = numfmt
        if pos:
            d.dLblPos = pos
        return d

    def style_chart(ch, title):
        ch.title, ch.style = title, 2
        ch.height, ch.width = 7.9, 12.2
        if ch.legend is not None:
            ch.legend.position = "b"
        try:
            if ch.y_axis is not None and ch.y_axis.majorGridlines is not None:
                ch.y_axis.majorGridlines.spPr = GraphicalProperties(
                    ln=LineProperties(solidFill=LINE, w=6350))
        except AttributeError:
            pass

    ch1 = BarChart(); ch1.type = "col"; ch1.gapWidth = 55
    ch1.add_data(Reference(S, min_col=2, min_row=39, max_row=42), titles_from_data=True)
    ch1.set_categories(Reference(S, min_col=1, min_row=40, max_row=42))
    ch1.dataLabels = dlabels(show_val=True, numfmt="0%", pos="outEnd")
    ch1.y_axis.numFmt = "0%"; ch1.y_axis.scaling.max = 1; ch1.y_axis.scaling.min = 0
    ch1.series[0].graphicalProperties.solidFill = IRON
    ch1.legend = None
    style_chart(ch1, "Branş Bazlı Ironman İlerlemesi")
    D.add_chart(ch1, "B34")

    ch2 = LineChart()
    ch2.add_data(Reference(S, min_col=3, min_row=18, max_row=30), titles_from_data=True)
    ch2.set_categories(Reference(S, min_col=2, min_row=19, max_row=30))
    s0 = ch2.series[0]
    s0.graphicalProperties.line.solidFill = IRON
    s0.graphicalProperties.line.width = 22000
    s0.marker = Marker(symbol="circle", size=6)
    s0.smooth = False
    ch2.legend = None
    style_chart(ch2, "Haftalık Antrenman Süresi (saat) — son 12 hafta")
    D.add_chart(ch2, "H34")

    ch3 = BarChart(); ch3.type = "col"; ch3.grouping = "stacked"
    ch3.overlap = 100; ch3.gapWidth = 45
    ch3.add_data(Reference(S, min_col=4, max_col=6, min_row=18, max_row=30), titles_from_data=True)
    ch3.set_categories(Reference(S, min_col=2, min_row=19, max_row=30))
    for srs, col in zip(ch3.series, (SWIM, BIKE, RUN)):
        srs.graphicalProperties.solidFill = col
        srs.graphicalProperties.line.solidFill = col
    style_chart(ch3, "Haftalık Toplam Mesafe (km) — branş dağılımı")
    D.add_chart(ch3, "B51")

    ch4 = PieChart()
    ch4.add_data(Reference(S, min_col=2, min_row=33, max_row=36), titles_from_data=True)
    ch4.set_categories(Reference(S, min_col=1, min_row=34, max_row=36))
    ch4.dataLabels = dlabels(show_pct=True)
    ch4.series[0].data_points = [DataPoint(idx=i, spPr=GraphicalProperties(solidFill=c))
                                 for i, c in enumerate((SWIM, BIKE, RUN))]
    style_chart(ch4, "Yüzme / Bisiklet / Koşu — Süre Dağılımı")
    D.add_chart(ch4, "H51")


# ==========================================================================
# NASIL KULLANILIR
# ==========================================================================
def _sheet_help(N, athlete, sw_t, bk_t, rn_t):
    N.sheet_view.showGridLines = False
    widths(N, {"A": 2.2, "B": 22, "C": 20, "D": 20, "E": 20, "F": 20, "G": 20,
               "H": 20, "I": 14, "J": 2.2})
    fill_range(N, 1, 1, 90, 10, BG)
    N.row_dimensions[2].height = 40
    merge_set(N, 2, 2, 2, 9, "📖  NASIL KULLANILIR?", f(20, True, WHITE),
              al("left", "center", indent=1), NAVY)
    N.row_dimensions[3].height = 22
    merge_set(N, 3, 2, 3, 9,
              "Bu dosya IRONMAN Takip uygulamasından dışa aktarıldı — formüller canlıdır.",
              f(9.5, False, "94A3B8"), al("left", "center", indent=1), NAVY)

    def h2(row, text, color=SLATE):
        N.row_dimensions[row].height = 26
        merge_set(N, row, 2, row, 9, text, f(11.5, True, WHITE),
                  al("left", "center", indent=1), color)

    def para(row, text, bold=False, color=INK, size=10, height=18):
        N.row_dimensions[row].height = height
        merge_set(N, row, 2, row, 9, text, f(size, bold, color),
                  al("left", "center", True, indent=1))

    h2(5, "🗂️  SAYFA REHBERİ")
    guide = [
        ("DASHBOARD", "Sadece bakılır", "Genel seviye, branş seviyeleri, toplam istatistikler, grafikler."),
        ("SEVİYELER", "Sadece bakılır", "0–6 arası 7 seviyenin branş bazlı kriterleri; tamamlanma otomatik."),
        ("ANTRENMAN KAYIT", "★ ANA VERİ", "Tüm antrenmanlar. Uygulamadan aktarıldı; elle de ekleyebilirsiniz."),
        ("HAFTALIK TAKİP", "Otomatik", "Hafta hafta özet + önceki haftaya göre % değişim."),
        ("BRANŞ İLERLEME", "Otomatik", "Branş bazlı detaylı analiz ve seviye motorunun çıktıları."),
        ("BRICK", "Sadece tarih", "Bisiklet→koşu antrenmanları; tarihler otomatik dolduruldu."),
        ("HEDEFLER", "Hedef tarihi", "Kilometre taşları; sadece tarihleri düzenleyin."),
        ("AYARLAR", "İsteğe bağlı", "Açılır liste seçenekleri ve seviye eşikleri."),
    ]
    header_row(N, 6, ["Sayfa", "Rolünüz", "Ne işe yarar?", "", "", "", "", ""], 2, 26)
    N.merge_cells(start_row=6, start_column=4, end_row=6, end_column=9)
    N.cell(row=6, column=4).value = "Ne işe yarar?"
    rr = 7
    for name, role, desc in guide:
        N.row_dimensions[rr].height = 20
        a = N.cell(row=rr, column=2, value=name)
        a.font, a.fill = f(10, True), fill(WHITE)
        a.alignment = al("left", "center", indent=1)
        b = N.cell(row=rr, column=3, value=role)
        b.font = f(9.5, True, IRON if role.startswith("★") else MUTED)
        b.fill, b.alignment = fill(WHITE), al("center")
        merge_set(N, rr, 4, rr, 9, desc, f(9.5, False, INK), al("left", "center", indent=1), WHITE)
        apply_box(N, rr, 2, rr, 9)
        rr += 1

    h2(16, "🎯  SEVİYE SİSTEMİ NASIL ÇALIŞIR?")
    para(17, "Her branş için iki ayrı seviye hesaplanır:", bold=True)
    para(18, "•  MESAFE SEVİYESİ (Rozet):  En uzun tek antrenmanınıza bakar. Bir kez kazandınız mı asla düşmez.")
    para(19, "•  FORM SEVİYESİ (Süreklilik):  Son 30 gündeki toplam hacminize bakar. Ara verirseniz düşer.")
    para(20, "MEVCUT SEVİYE = bu ikisinin küçüğü. Sadece bir kez uzun mesafe yapmak yetmez; düzenli antrenman da gerekir.",
         bold=True, color=IRON)
    para(21, "GENEL IRONMAN SEVİYESİ = üç branşın en düşüğü. Triatlonda zincir en zayıf halkası kadar güçlüdür.")

    header_row(N, 23, ["Seviye", "Ad", "🏊 Yüzme", "", "🚴 Bisiklet", "", "🏃 Koşu", ""], 2, 26)
    N.merge_cells("D23:E23"); N.cell(row=23, column=4).value = "🏊 Yüzme"
    N.merge_cells("F23:G23"); N.cell(row=23, column=6).value = "🚴 Bisiklet"
    N.merge_cells("H23:I23"); N.cell(row=23, column=8).value = "🏃 Koşu"
    rr = 24
    for lvl in range(7):
        N.row_dimensions[rr].height = 20
        N.cell(row=rr, column=2, value=lvl).font = f(11, True, IRON)
        N.cell(row=rr, column=2).alignment = al("center")
        name = LEVEL_NAMES[lvl] + (" 🏆" if lvl == 6 else "")
        N.cell(row=rr, column=3, value=name).font = f(10, True)
        N.cell(row=rr, column=3).alignment = al("left", "center", indent=1)
        merge_set(N, rr, 4, rr, 5, LEVEL_DETAIL["swim"][lvl][0], f(10, False, SWIM), al("center"))
        merge_set(N, rr, 6, rr, 7, LEVEL_DETAIL["bike"][lvl][0], f(10, False, BIKE), al("center"))
        merge_set(N, rr, 8, rr, 9, LEVEL_DETAIL["run"][lvl][0], f(10, False, RUN), al("center"))
        fill_range(N, rr, 2, rr, 9, WHITE)
        apply_box(N, rr, 2, rr, 9)
        rr += 1

    h2(32, "✍️  ELLE ANTRENMAN EKLEMEK")
    steps = [
        "1)  ANTRENMAN KAYIT sayfasında ilk boş satıra geçin.",
        "2)  Tarih → gg.aa.yyyy   ·   Branş → açılır listeden seçin (🏊 / 🚴 / 🏃).",
        "3)  Mesafe → HER ZAMAN KİLOMETRE.  Yüzme 750 m = 0,75  ·  3.800 m = 3,8",
        "4)  Süre → HER ZAMAN DAKİKA.  1 saat 12 dk = 72  ·  3 saat 30 dk = 210",
        "5)  Bisikletin hemen ardından koştuysanız her iki satırda da «Brick mi? = Evet» seçin.",
        "6)  Pace, hız ve tüm özetler otomatik hesaplanır — sarı hücreler dışına yazmayın.",
    ]
    rr = 33
    for s in steps:
        para(rr, s)
        rr += 1
    para(40, "⚠️  Bu dosyada yaptığınız değişiklikler uygulamaya geri yazılmaz. "
             "Uygulama ile Excel'i karıştırmayın: veri kaynağınız uygulama olsun, Excel'i rapor/yedek olarak kullanın.",
         color=IRON, bold=True, height=30)
    para(42, f"Bol şans {athlete} — her hafta bir adım yeter. 🏁", bold=True, color=IRON,
         size=11, height=26)
